# -*- coding: utf-8 -*-
"""从《安全权益履约周报数据（9月）.xlsx》重新生成履约弹窗所需的 5 个明细 JSON。
只重生成明细（meeting_reject / inspection_detail / disinfect_data / equipment_detail / health_violation），
不动 compliance.json / accident.json（它们由 _gen_from_excel.py 单独维护）。

用法: py -3 _gen_details_from_excel.py "<xlsx 路径>"
"""
import sys, os, json, datetime, re

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL = sys.argv[1] if len(sys.argv) > 1 else r"C:\Users\牛艳朝\Desktop\安全权益履约周报数据（9月） (1).xlsx"
DATA_DIR = os.path.join(BASE_DIR, "data")

from openpyxl import load_workbook

SERIAL_EPOCH = datetime.datetime(1899, 12, 30)
NULL_TOKENS = ("NULL", "null")
ERR_TOKENS = ("#DIV/0!", "#REF!", "#VALUE!", "#N/A", "#NAME?", "#NUM!")

def is_dt_str(v):
    return isinstance(v, str) and len(v) >= 19 and v[4] == "-" and v[7] == "-" and (v[10] == " " or v[10] == "T")

def parse_dt(v):
    if isinstance(v, datetime.datetime):
        return v
    if isinstance(v, datetime.date):
        return datetime.datetime(v.year, v.month, v.day)
    if is_dt_str(v):
        s = v.replace("T", " ")[:19]
        try:
            return datetime.datetime.strptime(s, "%Y-%m-%d %H:%M:%S")
        except ValueError:
            return None
    return None

def to_serial(v):
    dt = parse_dt(v)
    if dt is None:
        return None
    return int((dt - SERIAL_EPOCH).total_seconds() // 86400)

def norm_value(v):
    if isinstance(v, str):
        s = v.strip()
        if s in NULL_TOKENS or s in ERR_TOKENS:
            return None, True
    return v, False

def is_time_val(v):
    return isinstance(v, datetime.time)

def time_to_str(v):
    return v.strftime("%H:%M:%S")

def finish(v, null_style):
    if v is None:
        return None if null_style == "null" else ""
    return v

def transform_header_cells(row, date_mode="serial"):
    out = []
    for v in row:
        if v is None:
            out.append(None)
        elif is_time_val(v):
            out.append(time_to_str(v))
        elif isinstance(v, (datetime.datetime, datetime.date)) or is_dt_str(v):
            out.append(to_serial(v) if date_mode == "serial" else "")
        else:
            out.append(v)
    return out

def transform_data_row(row, datetime_cols=None):
    out = []
    for idx, v in enumerate(row):
        v, is_placeholder = norm_value(v)
        if is_placeholder:
            out.append(None)
            continue
        if v is None:
            out.append("")
            continue
        if is_time_val(v):
            out.append(time_to_str(v))
            continue
        if datetime_cols and idx in datetime_cols:
            if isinstance(v, (datetime.datetime, datetime.date)) or is_dt_str(v):
                out.append(to_serial(v))
            else:
                out.append(v)
            continue
        if isinstance(v, (datetime.datetime, datetime.date)) or is_dt_str(v):
            dt = parse_dt(v)
            out.append(dt.strftime("%Y-%m-%d %H:%M:%S") if dt else str(v))
            continue
        if isinstance(v, float) and v.is_integer():
            out.append(int(v))
        else:
            out.append(v)
    return out

def _core(name):
    """剥掉 sheet 名里的周次/日期噪声，取核心关键词"""
    s = re.sub(r"^\d+年\d+周", "", str(name or ""))
    return s.strip()

def resolve_sheet(wb, keyword):
    """按关键词在 sheet 名里唯一命中"""
    cands = [s for s in wb.sheetnames if keyword in s or keyword in _core(s)]
    if len(cands) == 1:
        return cands[0]
    raise KeyError("sheet 关键词 %r 命中 %d 个: %s (all=%s)" % (keyword, len(cands), cands, wb.sheetnames))

def read_rows(ws, max_cols):
    rows = []
    for row in ws.iter_rows(values_only=True):
        vals = list(row)[:max_cols]
        vals += [None] * (max_cols - len(vals))
        rows.append(vals)
    return rows

def is_blank(r, key_cols=None):
    if key_cols:
        return all(r[i] is None or (isinstance(r[i], str) and r[i].strip() == "") for i in key_cols if i < len(r))
    return all(v is None or (isinstance(v, str) and v.strip() == "") for v in r)

# (输出文件, sheet 关键词, 保留列数, datetime 列, 关键身份列——全空则视为幻影行丢弃)
DETAIL_SPECS = [
    ("meeting_reject.json",   "早会-驳回清单",   33, {1, 2},    {0, 7}),
    ("inspection_detail.json","子项检核明细",    29, {26, 27},  {0, 7}),
    ("disinfect_data.json",   "餐箱消毒",        12, None,      {0}),
    ("equipment_detail.json", "抽检不合规",      25, None,      {0, 12}),
    ("health_violation.json", "健康证违规报表",  12, None,      {0, 4}),
]

def fix_health_expired_inspection(data):
    """健康证/回执单过期 处罚金额按截图标准重算（2026-09-28 用户口径，线下站点检核处罚）：
    证件不符 N元/人次 —— 人数（本站此类驳回人次）≤4→300 / ≤9→400 / >9→500。
    站长行写「罚款」列，督导行（区域督导/总部员工）写「督导罚款」列。返回改动行数。"""
    if not data or len(data) < 2:
        return 0
    hdr = data[0]
    def col(name):
        for i, h in enumerate(hdr):
            if str(h or "").strip() == name:
                return i
        return -1
    i_item, i_reason = col("检核子项"), col("质检原因")
    i_site, i_role = col("站点名称"), col("提交人岗位")
    i_fine, i_sfine = col("罚款"), col("督导罚款")
    if min(i_item, i_reason, i_site, i_role, i_fine, i_sfine) < 0:
        return 0
    def is_hit(r):
        item = str(r[i_item] or "")
        reason = str(r[i_reason] or "")
        return (("健康证" in item) or ("健康证" in reason)) and \
               any(k in reason for k in ("回执单过期", "过期", "证件不符"))
    hits = [r for r in data[1:] if is_hit(r)]
    cnt = {}
    for r in hits:
        k = str(r[i_site] or "").strip()
        cnt[k] = cnt.get(k, 0) + 1
    changed = 0
    for r in hits:
        k = str(r[i_site] or "").strip()
        n = 300 if cnt[k] <= 4 else (400 if cnt[k] <= 9 else 500)
        role = str(r[i_role] or "")
        tgt = i_fine if role == "站长" else i_sfine
        if r[tgt] != n:
            r[tgt] = n
            changed += 1
    return changed

def main():
    wb = load_workbook(EXCEL, read_only=True, data_only=True)
    report = {}
    for jname, keyword, max_cols, dt_cols, key_cols in DETAIL_SPECS:
        sheet = resolve_sheet(wb, keyword)
        rows = read_rows(wb[sheet], max_cols)
        data = [transform_header_cells(rows[0], "serial")]
        kept = dropped = 0
        for r in rows[1:]:
            if is_blank(r, key_cols):  # 丢掉导出工具产生的幻影空行（关键身份列全空）
                dropped += 1
                continue
            data.append(transform_data_row(r, dt_cols))
            kept += 1
        fixed = 0
        if jname == "inspection_detail.json":
            fixed = fix_health_expired_inspection(data)
        path = os.path.join(DATA_DIR, jname)
        with open(path, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        report[jname] = {"sheet": sheet, "rows": len(data), "kept": kept, "dropped_blank": dropped}
        print(jname, "OK <-", sheet, "| kept", kept, "| dropped_blank", dropped,
              ("| 健康证过期罚金修正 %d 行" % fixed) if fixed else "")
    wb.close()
    print("DONE", json.dumps(report, ensure_ascii=False))

if __name__ == "__main__":
    main()
