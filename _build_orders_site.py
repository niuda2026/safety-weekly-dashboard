# -*- coding: utf-8 -*-
r"""
安全事故看板「上周完成单量」站维度取数脚本（2026-09-27 固化）
源：M:\下载\_数据明细_<起>-<止>(日期)_*.xlsx（用户每周放一份，文件名带日期区间）
  - sheet 名任意（当前为 '0'），表头行含：配送站点名称 / 配送完成运单量（/ 加盟商名称 / 配送站点）
  - 配送站点名称 与 安全事故表站点行「区域」名一致（含 集约配送- 前缀变体）
  - 单元格可能出现字符串 'NULL' → 按 0 计
产出：data/orders_site.json
  { meta: {source_file, date_range, generated_at, station_count, total_orders},
    sites: { 站点名: 完成单量, ... } }
用法：
  python _build_orders_site.py            # 扫 M 盘最新一份
  python _build_orders_site.py --src <xlsx路径>
"""
import glob
import json
import os
import re
import sys
from datetime import datetime

import openpyxl

PROJ = os.path.dirname(os.path.abspath(__file__))
M_DIR = r"M:\下载"
OUT = os.path.join(PROJ, "data", "orders_site.json")


def fv(x):
    """单元格数值化：'NULL'/空 → 0.0"""
    if x is None:
        return 0.0
    s = str(x).strip()
    if s in ("", "NULL", "null", "-"):
        return 0.0
    try:
        return float(s)
    except ValueError:
        return 0.0


def pick_source():
    """M 盘扫 _数据明细_*.xlsx，取 mtime 最新；返回 (path, date_range)"""
    cands = glob.glob(os.path.join(M_DIR, "_数据明细_*.xlsx"))
    if not cands:
        return None, None
    cands.sort(key=os.path.getmtime, reverse=True)
    path = cands[0]
    m = re.search(r"(\d{4}-\d{2}-\d{2})-(\d{4}-\d{2}-\d{2})", os.path.basename(path))
    rng = (m.group(1) + "~" + m.group(2)) if m else ""
    return path, rng


def build(src):
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(c or "").strip() for c in rows[0]]
    try:
        i_site = hdr.index("配送站点名称")
        i_ord = hdr.index("配送完成运单量")
    except ValueError:
        raise SystemExit("表头缺「配送站点名称/配送完成运单量」列，实际表头：%s" % hdr)
    sites = {}
    null_rows = 0
    for r in rows[1:]:
        if not r or not r[i_site]:
            continue
        name = str(r[i_site]).strip()
        if str(r[i_ord]).strip() == "NULL":
            null_rows += 1
        sites[name] = sites.get(name, 0.0) + fv(r[i_ord])
    total = sum(sites.values())
    return sites, total, null_rows


def main():
    src = None
    if "--src" in sys.argv:
        src = sys.argv[sys.argv.index("--src") + 1]
        m = re.search(r"(\d{4}-\d{2}-\d{2})-(\d{4}-\d{2}-\d{2})", os.path.basename(src))
        rng = (m.group(1) + "~" + m.group(2)) if m else ""
    else:
        src, rng = pick_source()
    if not src or not os.path.isfile(src):
        raise SystemExit("M 盘未找到 _数据明细_*.xlsx 源文件")
    sites, total, null_rows = build(src)
    data = {
        "meta": {
            "source_file": os.path.basename(src),
            "date_range": rng,
            "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "station_count": len(sites),
            "total_orders": int(total),
            "null_order_rows": null_rows,
        },
        "sites": sites,
    }
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False)
    print("OK %s" % OUT)
    print("  源文件: %s（%s）" % (os.path.basename(src), rng))
    print("  站点数: %d  总单量: %d  NULL行: %d" % (len(sites), int(total), null_rows))


if __name__ == "__main__":
    main()
